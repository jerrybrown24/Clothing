"""Export every input data point used in the Circle Mall feasibility report (plus key derived
outputs) to Circle_Mall_Inputs.xlsx. Run after forecast.py and pnl.py."""
import json
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

R = json.load(open("results.json"))
SRC_COG = "COG_month-loc_salesfootfall_from_2024.xlsx"
SRC_PL = "Circle_Mall_002.xlsx"
sheets = {}

# 1. Store & lease setup ------------------------------------------------------------------
sheets["1 Store & lease setup"] = pd.DataFrame([
    ["Store", "Adults - Circle Mall (Cotton On, LOC_ID C49)", "", "P&L workbook / COG file"],
    ["Location", "Circle Mall, Jumeirah Village Circle, Dubai", "", ""],
    ["GLA", 3598, "sq ft", "P&L sheet 'P & L' (Total Sqft)"],
    ["Lease anniversary", "8 Aug (lease year 8 Aug - 7 Aug)", "", "Sheet1 (2) start/end dates"],
    ["Old rent basis", 0.14, "% of sales (turnover rent)", "Sheet1 (2) Rent %; P&L FY'25 rent / sales"],
    ["Scenario 1 rent, years 1 / 2 / 3", "246 / 290 / 304", "AED / sq ft / yr gross", "User instruction"],
    ["Scenario 2 rent, years 1 / 2 / 3", "276 / 290 / 304", "AED / sq ft / yr gross", "User instruction"],
    ["Service charge (held flat)", 41, "AED / sq ft / yr", "P&L sheet 'P & L' (Sc/ Sq ft)"],
    ["Marketing levy (held flat)", 15, "AED / sq ft / yr", "P&L sheet 'P & L' (MKT/Sqft)"],
    ["Renewal start", "8 Aug 2026", "", "User instruction"],
    ["Horizon", 3, "lease years (to 7 Aug 2029)", "User instruction"],
    ["Sales year 1", "Model forecast", "weighted ensemble of 8 methods", "Section 2"],
    ["Sales year 2", 0.0, "no like-for-like growth (Al Khail Avenue mall opening)", "User instruction"],
    ["Sales year 3", -0.10, "10% decline on year 2 (Al Khail Avenue mall)", "User instruction"],
    ["Turnover rent clause", 0.12, "higher of base rent or 12% of sales", "Scenario / Circle Mall Rent sheets (TOR%)"],
    ["Tax rate used for PAT", 0.09, "UAE corporate tax, store level", "Assumption (UAE CT 9%)"],
    ["Data cut-off", "Aug-2026", "Sep-26 excluded (partial month)", "COG file"],
    ["Market rent comparables", "230 - 240", "AED / sq ft gross", "User-verified: Apparel Group new rentals at Circle Mall"],
    ["Mall annual footfall", 7_500_000, "visitors / yr", "P&L sheet 'notes'"],
    ["Healthy rent-to-sales ceiling", 0.20, "", "Report assumption (apparel benchmark)"],
    ["Healthy rent-to-income ceiling", 0.32, "", "Report assumption"],
], columns=["Item", "Value", "Unit / note", "Source"])

# 2. Rent & sales history ------------------------------------------------------------------
sheets["2 Rent history"] = pd.DataFrame([
    ["8 Aug 2022 - 7 Aug 2023", 3345376, 468352.64, 0.14, 130.17],
    ["8 Aug 2023 - 7 Aug 2024", 3476676, 486734.64, 0.14, 135.28],
    ["8 Aug 2024 - 7 Aug 2025", 3375376, 472552.64, 0.14, 131.34],
    ["8 Aug 2025 - 7 Aug 2026", 3778830.96, 529036.33, 0.14, 147.04],
], columns=["Lease year", "Sales (AED)", "Rent (AED)", "Rent %", "Rent / sq ft"]).assign(Source="P&L sheet 'Sheet1 (2)'")
sheets["2b Lease-year sales (notes)"] = pd.DataFrame(
    [[k, v] for k, v in R["hist_ly"].items()], columns=["Lease year", "Sales (AED)"]).assign(Source="P&L sheet 'notes'")

# 3. Benchmarks ---------------------------------------------------------------------------
sheets["3 Benchmarks"] = pd.DataFrame([
    [b["store"], b["sqft"], round(b["psf"], 2), round(b["rent"]), round(b["ttm"]), round(b["sales_psf"]), round(b["rts"], 4)]
    for b in R["bench"]], columns=["Store", "Sq ft", "Gross rent / sq ft", "Annual rent", "TTM sales Sep-25..Aug-26",
                                   "Sales / sq ft", "Rent-to-sales"]).assign(
    Source="Size & rent: P&L 'notes' (Dubai Hills 244/65/20/6.10; Dubai Mall 365/80/20/9.13; Mirdif 153.18/69.81/17.38/21.41); sales: COG file")

# 4-6. P&L cost inputs -----------------------------------------------------------------------
lines = ["Sales", "Gross Margin", "Net Income", "Employee expense", "Rent", "Advertising & marketing",
         "Other opex", "Shared common", "Shared others", "Depreciation", "Finance & extraordinary"]
fy25 = [3635846.29, 2505247.92, 2235204.25, 246033.86, 509018.48, 33418.94, 278611.63, 183849.07, 86106.41, 259318.77, 0]
augdec = [2209715.49, 1635655.72, 1415861.99, 151647.0, 368795, 31481.69, 110781.10, 74543.30, 30687.36, 0, 1115.0]
sheets["4 P&L FY25 actual"] = pd.DataFrame({"Line": lines, "FY'25 (AED)": fy25,
                                            "Source": "P&L sheet 'P & L', column Fy'25"})
ws = openpyxl.load_workbook(SRC_PL, data_only=True)["YR'26"]
rows = {r[2]: list(r[3:10]) for r in ws.iter_rows(min_row=3, max_row=17, values_only=True)}
keymap = [("Sales", "Sales"), ("Gross Margin", "Gross Margin"), ("Net Income", "Net Income"), ("emp.Exp", "Employee expense"),
          ("Rent", "Rent"), ("Adv &Mktg Exp", "Advertising & marketing"), ("Other Admin", "Other opex"),
          ("Shared Cost Emp", "Shared common"), ("Shared Cost Others", "Shared others"), ("Depreciation", "Depreciation"),
          ("Finance & Extrdnry", "Finance & extraordinary")]
y26 = pd.DataFrame([[lab] + rows[k] for k, lab in keymap],
                   columns=["Line", "Jan-26", "Feb-26", "Mar-26", "Apr-26", "May-26", "Jun-26", "Jul-26"])
y26["Jan-Jul total"] = y26.iloc[:, 1:8].sum(axis=1)
y26["Source"] = "P&L sheet \"YR'26\" monthly actuals"
sheets["5 P&L Jan-Jul 26 actual"] = y26
sheets["6 P&L Aug-Dec 26 estimate"] = pd.DataFrame({"Line": lines, "Aug-Dec 2026 (AED)": augdec,
    "Source": "P&L sheet 'P & L', column 'Aug to DEC' (sales/margins replaced by forecast in the report; costs used as given)"})

fwd = pd.DataFrame({
    "Line": ["Employee expense", "Advertising & marketing", "Other opex", "Shared common", "Shared others", "Depreciation", "Finance & extraordinary"],
    "Yr'27": [382176.05, 91379.44, 280819.93, 203661.90, 93040.68, 15213.98, 4878.21],
    "Yr'28": [389819.57, 93207.03, 286436.33, 207735.13, 94901.50, 0, 0],
    "Yr'29": [397615.96, 95071.17, 292165.06, 211889.84, 96799.53, 0, 0],
    "Yr'30": [405568.28, 96972.60, 298008.36, 216127.63, 98735.52, 0, 0],
    "Yr'31": [413679.65, 98912.05, 303968.53, 220450.18, 100710.23, 0, 0]})
fwd = pd.concat([fwd, pd.DataFrame([["Gross margin % of sales", *[R["consts"]["gm_pct"]] * 5],
                                    ["Net Income % of sales", *[R["consts"]["ni_pct"]] * 5]], columns=fwd.columns)])
fwd["Source"] = "P&L sheet 'P & L 5 Yrs' (GM% and NI% from Yr'27: 3,377,687 / 4,564,441 and 2,921,242 / 4,564,441)"
sheets["7 P&L forward costs"] = fwd

# 8. Monthly sales & footfall --------------------------------------------------------------
s = pd.read_excel(SRC_COG, sheet_name="sales", header=1)
f = pd.read_excel(SRC_COG, sheet_name="footfall", header=1)
cm = pd.DataFrame({"Month": R["months"], "Circle Mall net sales (AED)": R["actual"], "Circle Mall footfall": R["footfall"]})
cm["Sales per visitor"] = cm.iloc[:, 1] / cm.iloc[:, 2]
cm["Source"] = "COG file, sheets 'sales' & 'footfall' (BRAND_ID COT, LOC_ID C49)"
sheets["8 Circle Mall monthly"] = cm

# 9. Peer stores ---------------------------------------------------------------------------
c = s[(s.BRAND_ID == "COT") & (s.MONTH_ID <= 202608)]
peer = c.pivot_table(index="MONTH_ID", columns="STORE_LOCATION", values="NET_SALES", aggfunc="sum").reset_index()
peer.insert(1, "Like-for-like chain total (peers used)", peer[R["peers"]].sum(axis=1))
sheets["9 Cotton On stores monthly"] = peer
sheets["9b Peer growth Jan-Aug"] = pd.DataFrame([[k, v] for k, v in R["peer_yoy"].items()],
    columns=["Store", "Sales growth Jan-Aug 26 vs Jan-Aug 25"]).assign(
    Note="Like-for-like peers = stores trading every month: " + ", ".join(R["peers"]))
ff = f[(f.BRAND_ID == "COT") & (f.MONTH_ID <= 202608)].pivot_table(index="MONTH_ID", columns="STORE_LOCATION", values="FOOTFALL", aggfunc="sum").reset_index()
sheets["9c Cotton On footfall monthly"] = ff

# 10. Online channels -----------------------------------------------------------------------
on = s[s.STORE_LOCATION.isin(["COG Online Store", "Namshi Cotton On Marketplace"]) & (s.MONTH_ID <= 202608)]
sheets["10 Online channels"] = on.pivot_table(index="MONTH_ID", columns="STORE_LOCATION", values="NET_SALES").reset_index().assign(Source="COG file 'sales'")

# 11. External drivers --------------------------------------------------------------------
drv = [[d["name"], d["value"] / 100, d["weight"], d["value"] * d["weight"] / 100, d["note"]] for d in R["drivers"]]
drv += [[a["name"], a["value"] / 100, "", a["value"] / 100, "Deduction (analyst judgement)"] for a in R["adjust"]]
drv += [["External growth rate (total)", R["macro_g"] / 100, "", R["macro_g"] / 100, "Reference only: not used for years 2-3, which follow the Al Khail Avenue assumption"]]
sheets["11 External growth drivers"] = pd.DataFrame(drv, columns=["Driver", "Growth p.a.", "Weight", "Contribution", "Basis / source"])

# 12-14. Derived outputs --------------------------------------------------------------------
mods = [[k, v["family"], v["mape"] / 100, v["weight"], *v["lease_years"]] for k, v in R["models"].items()]
mods.append(["Weighted ensemble", "", R["ensemble"]["mape"] / 100, 1.0, *R["ensemble"]["lease_years"]])
sheets["12 Model backtest (derived)"] = pd.DataFrame(mods, columns=["Method", "Family", "Backtest MAPE Jan-Aug 26", "Ensemble weight",
                                                                   "LY 26-27", "LY 27-28", "LY 28-29"])
L0 = R["ly0"]
cols = ["sales", "gm", "ni", "Employee expense", "rent", "Advertising & marketing", "Other opex", "Shared common", "Shared others",
        "total_exp", "store_profit", "dep", "fin", "net", "rent_psf", "rent_to_sales", "rent_to_ni"]
names = ["Net sales", "Gross margin", "Net Income (Income)", "Employee expense", "Rent (gross)", "Advertising & marketing", "Other opex",
         "Shared common", "Shared others", "Total expense", "Store / cash profit", "Depreciation", "Finance", "Net profit (pre-tax)",
         "Rent / sq ft", "Rent-to-sales", "Rent-to-income"]
patf = lambda d: d["net"] * 0.91 if d["net"] > 0 else d["net"]
for sc in ("S1", "S2"):
    P = R["pnl"]["base"][sc]["s1"]
    ly = pd.DataFrame({"Line": names, "8 Aug 25 - 7 Aug 26 (actual, old rent)": [L0[c] for c in cols],
                       **{f"Year {i+1}: 8 Aug {26+i} - 7 Aug {27+i}": [P[i][c] for c in cols] for i in range(3)}})
    sheets[f"14{'ab'[sc == 'S2']} Lease-year P&L {sc} (derived)"] = pd.concat([ly, pd.DataFrame([["PAT (9% tax)", patf(L0), *[patf(p) for p in P]]], columns=ly.columns)])
sheets["13 Monthly forecast (derived)"] = pd.DataFrame({"Month": R["fut_months_ext"], "Adopted forecast net sales (AED)": R["forecast_monthly"]})
sheets["13b Lease-year sales (derived)"] = pd.DataFrame({"Lease year": ["8 Aug 26 - 7 Aug 27", "8 Aug 27 - 7 Aug 28", "8 Aug 28 - 7 Aug 29"],
    "Adopted sales": R["scenarios"]["base"][:3], "Model + market trend (reference)": R["scenarios"]["trend"][:3]})

# write ----------------------------------------------------------------------------------
out = "Circle_Mall_Inputs.xlsx"
with pd.ExcelWriter(out, engine="openpyxl") as xw:
    idx = pd.DataFrame([[n, ("Derived output" if "derived" in n else "Input")] for n in sheets], columns=["Sheet", "Type"])
    idx.to_excel(xw, sheet_name="Index", index=False)
    for n, df in sheets.items():
        df.to_excel(xw, sheet_name=n[:31], index=False)
wb = openpyxl.load_workbook(out)
hdr = PatternFill("solid", fgColor="0D5F52")
for ws in wb.worksheets:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = hdr; cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, col in enumerate(ws.columns, 1):
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), 60)
        for c in col[1:]:
            if isinstance(c.value, float):
                c.number_format = "0.0%" if 0 < abs(c.value) < 1.5 else "#,##0"   # rates are stored as fractions
    ws.freeze_panes = "B2"
wb.save(out)
print("wrote", out, len(sheets), "sheets")
